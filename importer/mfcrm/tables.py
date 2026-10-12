"""Turn a report file (csv, xls, xlsx) into tables of rows keyed by header name."""
import csv
import io
import os
import re
import shutil
import subprocess
import tempfile


def norm(header) -> str:
    """'PAN Number' -> 'pannumber', 'trxn_sign' -> 'trxnsign'."""
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


class Table:
    def __init__(self, sheet, headers, rows):
        self.sheet = sheet
        self.headers = [norm(h) for h in headers]
        self.original_headers = [str(h or "").strip() for h in headers]
        self.rows = rows            # list of dicts: normalised header -> value

    @property
    def columns(self):
        return set(self.headers)


def _find_header(rows):
    """Header = first row with at least five filled cells."""
    for i, row in enumerate(rows[:20]):
        if sum(1 for v in row if str(v if v is not None else "").strip()) >= 5:
            return i
    return None


def _to_table(sheet, rows):
    rows = [list(r) for r in rows]
    h = _find_header(rows)
    if h is None:
        return None
    headers = rows[h]
    keys = [norm(x) for x in headers]
    out = []
    for r in rows[h + 1:]:
        if not any(str(v if v is not None else "").strip() for v in r):
            continue
        d = {}
        for k, v in zip(keys, r):
            if k and k not in d:          # first column wins if a name repeats
                d[k] = v
        out.append(d)
    return Table(sheet, headers, out)


def _read_csv(data: bytes):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:5000]
    delim = "|" if sample.count("|") > sample.count(",") else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    t = _to_table("csv", rows)
    return [t] if t else []


def _read_xlsx(data: bytes):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        t = _to_table(ws.title, ws.iter_rows(values_only=True))
        if t:
            out.append(t)
    return out


def _read_xls(data: bytes):
    try:
        import xlrd
    except ImportError:
        return _read_xls_via_libreoffice(data)
    book = xlrd.open_workbook(file_contents=data)
    out = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value)
            rows.append(row)
        t = _to_table(sh.name, rows)
        if t:
            out.append(t)
    return out


def _read_xls_via_libreoffice(data: bytes):
    """Fallback for machines without xlrd: convert with LibreOffice."""
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        raise RuntimeError("Reading .xls needs the xlrd package or LibreOffice")
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "in.xls")
        with open(src, "wb") as f:
            f.write(data)
        subprocess.run([office, "--headless", "--convert-to", "xlsx", "--outdir", tmp, src],
                       check=True, capture_output=True, timeout=180)
        with open(os.path.join(tmp, "in.xlsx"), "rb") as f:
            return _read_xlsx(f.read())


def read_tables(name: str, data: bytes):
    ext = os.path.splitext(name.lower())[1]
    if data[:4] == b"PK\x03\x04" and ext not in (".zip",):
        return _read_xlsx(data)
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return _read_xls(data)
    if ext in (".csv", ".txt") or ext == "":
        return _read_csv(data)
    if ext == ".xlsx":
        return _read_xlsx(data)
    if ext == ".xls":
        # Some ".xls" files from registrars are really HTML or text tables.
        return _read_csv(data)
    raise ValueError(f"Unsupported file type: {name}")
